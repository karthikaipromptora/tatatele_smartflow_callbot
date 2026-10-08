"""Read call lists from .xlsx / .csv uploads (our template or the Smartflo collections export) and build the template."""
import csv
import io
import re
from datetime import date, datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from helpers.call_input import email_domain, normalize_amount, normalize_phone

MAX_ROWS = 500
MAX_BYTES = 5 * 1024 * 1024

# Our template: canonical field → (header shown in the template, accepted header spellings)
COLUMNS = {
    "phone_number":   ("Phone number",   ["phone", "phonenumber", "mobile", "mobilenumber", "number", "contact", "contactnumber", "customerphone", "customernumber"]),
    "customer_name":  ("Customer name",  ["name", "customer", "customername", "company", "companyname", "account", "accountname"]),
    "amount":         ("Amount due",     ["amount", "amountdue", "dueamount", "outstanding", "outstandingamount", "pendingamount", "balance"]),
    "billing_period": ("Billing period", ["billingperiod", "period", "billingmonth", "month", "billperiod", "invoicemonth"]),
    "due_date":       ("Due date",       ["duedate", "paymentduedate", "due"]),
    "invoice_number": ("Invoice number", ["invoice", "invoicenumber", "invoiceno", "billno", "billnumber"]),
    "service_name":   ("Service",        ["service", "servicename", "product"]),
    "voice_id":       ("Voice",          ["voice", "voiceid"]),
}
REQUIRED = ["phone_number", "customer_name", "amount", "billing_period"]

# Columns of the Smartflo "Invoice wise" export that identify it.
SMARTFLO_SIGNATURE = {"BILL_REF_NO", "BALANCE_DUE_OVER_INVOICE", "PAYMENT_DUE_DATE", "BILL_COMPANY"}
# Columns offered as pre-call filters, with display labels.
FILTER_COLUMNS = {"region": ("Region", "REGION"), "partner": ("Partner", "Partner Name"),
                  "product": ("Product", "ACCOUNT_PRODUCT"), "bucket": ("Ageing bucket", "BUCKET_NAME")}


class BatchFileError(ValueError):
    pass


def _key(header) -> str:
    return re.sub(r"[^a-z0-9]", "", str(header or "").lower())


_ALIASES = {alias: field for field, (_, aliases) in COLUMNS.items() for alias in aliases}


def _cell_text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.date().isoformat() if v.time() == datetime.min.time() else v.isoformat(sep=" ")
    if isinstance(v, date):
        return v.isoformat()
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _non_empty(rows):
    return [r for r in rows if any(str(c).strip() for c in r if c is not None)]


def _read_tables(filename: str, data: bytes) -> dict[str, list[list]]:
    """All sheets as {sheet name: rows}; a CSV becomes a single sheet."""
    if len(data) > MAX_BYTES:
        raise BatchFileError("File is larger than 5 MB")
    name = filename.lower()
    if name.endswith(".csv"):
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = data.decode("latin-1")
        return {"CSV": _non_empty(csv.reader(io.StringIO(text)))}
    if name.endswith(".xlsx"):
        try:
            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        except Exception as e:
            raise BatchFileError("This file could not be read as an Excel workbook") from e
        try:
            return {ws.title: _non_empty(list(r) for r in ws.iter_rows(values_only=True)) for ws in wb.worksheets}
        finally:
            wb.close()
    raise BatchFileError("Upload an Excel (.xlsx) or CSV (.csv) file")


_TRUNCATED_WORDS = ("LIMITED", "PRIVATE", "SERVICES", "TECHNOLOGIES", "INDUSTRIES", "SOLUTIONS", "ENTERPRISES")


def _company_name(raw: str) -> str:
    """'KRISHNA INSTITUTE OF MEDICAL SCIENCES LI' → 'Krishna Institute of Medical Sciences Limited'.

    The billing export cuts names at 40 characters; a cut-off final word is completed, and all-caps
    names are converted to title case because TTS spells out all-caps words letter by letter.
    """
    text = re.sub(r"\s+", " ", str(raw or "")).strip()
    if len(str(raw or "").rstrip()) >= 40 and " " in text:
        head, last = text.rsplit(" ", 1)
        full = next((w for w in _TRUNCATED_WORDS if w.startswith(last.upper()) and len(last) >= 2), None)
        if full:
            text = f"{head} {full if text.isupper() else full.capitalize()}"
    if not text or not text.isupper():
        return text
    small = {"of", "and", "the", "for", "in", "on", "at", "&"}
    keep_upper = {"IT", "BPO", "LLP", "HDFC", "ICICI", "SBI", "TCS", "HCL", "IBM", "NTPC", "ONGC", "GAIL"}
    words = []
    for i, w in enumerate(text.split(" ")):
        if w in keep_upper:
            words.append(w)
        elif i and w.lower() in small:
            words.append(w.lower())
        else:
            words.append(w.capitalize())
    return " ".join(words)


def _rows_from_smartflo(tables: dict[str, list[list]]) -> list[dict]:
    invoice_sheet = next((rows for rows in tables.values() if rows and SMARTFLO_SIGNATURE <= {str(h).strip() for h in rows[0] if h}), None)
    header = [str(h).strip() if h is not None else "" for h in invoice_sheet[0]]

    # Net balance per account from the "Account Wise" sheet, if present.
    net_balance: dict[str, float] = {}
    for rows in tables.values():
        if rows and {"ACCOUNT_NO", "Net Bal"} <= {str(h).strip() for h in rows[0] if h}:
            h = [str(x).strip() if x is not None else "" for x in rows[0]]
            ai, ni = h.index("ACCOUNT_NO"), h.index("Net Bal")
            for r in rows[1:]:
                try:
                    net_balance[_cell_text(r[ai])] = float(r[ni])
                except (TypeError, ValueError, IndexError):
                    pass

    out = []
    for n, raw in enumerate(invoice_sheet[1:], start=2):
        rec = {}
        for i, h in enumerate(header):
            if h and i < len(raw) and h not in rec:  # first occurrence wins (ACCOUNT_STATUS appears twice)
                rec[h] = raw[i]
        account = _cell_text(rec.get("ACCOUNT_NO"))
        phone = _cell_text(rec.get("CUST_PHONE1"))
        if not normalize_phone(phone) and normalize_phone(_cell_text(rec.get("CUST_PHONE2"))):
            phone = _cell_text(rec.get("CUST_PHONE2"))
        try:
            balance = round(float(rec.get("BALANCE_DUE_OVER_INVOICE") or 0))
        except (TypeError, ValueError):
            balance = 0
        try:
            paid = round(float(rec.get("AMT_PAID_AGAINST_INVOICE") or 0))
        except (TypeError, ValueError):
            paid = 0
        statement = rec.get("STATEMENT_DATE")

        skip = None
        status = _cell_text(rec.get("ACCOUNT_STATUS")).upper()
        if status and status != "ACTIVE":
            skip = f"Account is {status.lower()}"
        elif account in net_balance and net_balance[account] <= 0:
            credit = normalize_amount(abs(net_balance[account]))
            skip = f"Account has a credit balance of ₹{credit}" if credit else "Account has no net balance due"

        out.append({
            "row": n,
            "phone_number": phone,
            "customer_name": _company_name(_cell_text(rec.get("BILL_COMPANY"))),
            "amount": str(balance) if balance > 0 else "",
            "billing_period": statement if isinstance(statement, (date, datetime)) else _cell_text(statement),
            "due_date": rec.get("PAYMENT_DUE_DATE"),
            "invoice_number": _cell_text(rec.get("BILL_REF_NO")),
            "account_number": account,
            "service_name": _cell_text(rec.get("ACCOUNT_PRODUCT")),
            "amount_paid": str(paid) if paid > 0 else "",
            "email_domain": email_domain(rec.get("CUST_EMAIL")),
            "_skip": skip,
            "_filters": {k: _cell_text(rec.get(col)) for k, (_, col) in FILTER_COLUMNS.items()},
            "_source": {h: _cell_text(v) for h, v in rec.items() if _cell_text(v)},
        })
    return out


def _rows_from_template(table: list[list]) -> list[dict]:
    mapping = {}
    for col, header in enumerate(table[0]):
        field = _ALIASES.get(_key(header))
        if field and field not in mapping.values():
            mapping[col] = field
    missing = [COLUMNS[f][0] for f in REQUIRED if f not in mapping.values()]
    if missing:
        raise BatchFileError(f"Missing column{'s' if len(missing) > 1 else ''}: {', '.join(missing)}. Use the template for the expected headers.")
    headers = [_cell_text(h) for h in table[0]]
    rows = []
    for n, raw in enumerate(table[1:], start=2):
        values = {field: (raw[col] if col < len(raw) else None) for col, field in mapping.items()}
        if not any(_cell_text(v) for v in values.values()):
            continue
        rows.append({
            "row": n, **values, "_skip": None, "_filters": {},
            "_source": {h: _cell_text(raw[i]) for i, h in enumerate(headers) if h and i < len(raw) and _cell_text(raw[i])},
        })
    return rows


def parse_rows(filename: str, data: bytes) -> tuple[str, list[dict]]:
    """Return (format, rows). format is 'smartflo' or 'template'."""
    tables = {name: rows for name, rows in _read_tables(filename, data).items() if rows}
    if not tables:
        raise BatchFileError("The file is empty")
    if any(SMARTFLO_SIGNATURE <= {str(h).strip() for h in rows[0] if h} for rows in tables.values()):
        fmt, rows = "smartflo", _rows_from_smartflo(tables)
    else:
        fmt, rows = "template", _rows_from_template(next(iter(tables.values())))
    if not rows:
        raise BatchFileError("The file has headers but no customer rows")
    if len(rows) > MAX_ROWS:
        raise BatchFileError(f"The file has {len(rows)} rows; the limit is {MAX_ROWS} per upload")
    return fmt, rows


# The template mirrors the Smartflo collections export column for column, so a filled template and a
# raw export are read the same way. Only the columns marked "AI" reach the call; the rest are kept
# with the call for filters and the call log.
INVOICE_COLUMNS = [
    "CIRCLE", "ACCOUNT_NO", "ACCOUNT_CATEGORY", "MARKET_SEGMENT", "BILL_PERIOD", "SEGMENT", "REGION", "4 Region",
    "GTM", "Partner Name", "ACCOUNT_PRODUCT", "ACCOUNT_STATUS", "BILL_MODE", "CUST_EMAIL", "ACCOUNT_ACTIVATION_DATE",
    "BILL_REF_NO", "BILL_REF_RESETS", "BILL_SEQUENCE_NUM", "STATEMENT_DATE", "PREP_DATE", "PAYMENT_DUE_DATE",
    "AGEING_OF_INVOICE", "BUCKET_NAME", "BUCKET_NUMBER", "TOTAL_INVOICE_AMOUNT", "AMT_PAID_AGAINST_INVOICE",
    "LAST_PAYMENT_DT_TO_INVOICE", "ADJ_POSTED_AGAINST_INVOICE", "LAST_ADJ_DT_AGAINST_INVOICE",
    "BALANCE_DUE_OVER_INVOICE", "ACCOUNT_STATUS", "LAST_DISCONNECTION_TSP_DATE", "BUCKET_MONTH", "BILL_COMPANY",
    "FC_CODE", "REV_RCV_COST_CTR", "OWNING_COST_CTR", "COVID", "COMPANY_ID", "CUST_PHONE1", "CUST_PHONE2",
]
ACCOUNT_COLUMNS = [
    "ACCOUNT_NO", "Logo ID", "BILL_PERIOD", "SEGMENT", "Region", "4 Region", "STD NAME", "Product Type",
    "ACCOUNT_STATUS", "Bill Delivery Mode", "Email ID", "Alternate number 1", "Alternate number 2", "E", "PreDue",
    "0-30", "30-60", "60-90", "90-120", "120-150", "150-180", "180-240", "240-360", "360-720", ">720",
    "Total Bal", "Credit Balance", "Net Bal", "Aging Bucket", "Customer Address", "GTM", "Mapped Partner Name",
]
# Invoice-wise columns the bot uses: column → (role, required, example, what it is used for)
INVOICE_ROLES = {
    "CUST_PHONE1":              ("Dial", True,  "9876543210", "Number that is called (10 digits; +91 or a leading 0 is fine)."),
    "CUST_PHONE2":              ("Dial", False, "9876500000", "Called instead when CUST_PHONE1 is missing or invalid."),
    "BILL_COMPANY":             ("AI",   True,  "EXAMPLE INDUSTRIES PRIVATE LIMITED", "Customer name the agent addresses. Names cut at 40 characters are completed."),
    "BALANCE_DUE_OVER_INVOICE": ("AI",   True,  "25000", "Amount due on this invoice, in rupees. Rounded to whole rupees."),
    "STATEMENT_DATE":           ("AI",   True,  "03-09-2026", "Invoice date; spoken as the billing month (September 2026)."),
    "PAYMENT_DUE_DATE":         ("AI",   True,  "20-09-2026", "Due date. Decides overdue vs pre-due and the days overdue."),
    "BILL_REF_NO":              ("AI",   True,  "4846808977", "Invoice number, read out if the customer asks."),
    "ACCOUNT_NO":               ("AI",   False, "209319169", "Account number, read out if the customer asks. Links to the Account Wise sheet."),
    "ACCOUNT_PRODUCT":          ("AI",   False, "Managed Enterprise Internet Service", "Service name. Also a filter."),
    "AMT_PAID_AGAINST_INVOICE": ("AI",   False, "5000", "Part payment already received, acknowledged on the call."),
    "CUST_EMAIL":               ("AI",   False, "accounts@example.com", "Only the domain (example.com) is mentioned, when asked where the invoice was sent."),
    "ACCOUNT_STATUS":           ("Skip", False, "ACTIVE", "Rows whose account is not ACTIVE are skipped (first ACCOUNT_STATUS column)."),
    "REGION":                   ("Filter", False, "AP", "Filter before calling."),
    "Partner Name":             ("Filter", False, "KMS", "Filter before calling."),
    "BUCKET_NAME":              ("Filter", False, "0-30 Days", "Filter before calling."),
}
ACCOUNT_ROLES = {
    "ACCOUNT_NO": ("Skip", "Matches the invoice rows to this account."),
    "Net Bal":    ("Skip", "Invoices of accounts whose net balance is 0 or less (credit) are skipped."),
}
_ROLE_FILL = {"Dial": "0B6BCB", "AI": "002F51", "Skip": "C2571A", "Filter": "4A7A2A"}
_TEXT_COLUMNS = {"ACCOUNT_NO", "BILL_REF_NO", "CUST_PHONE1", "CUST_PHONE2", "COMPANY_ID", "Logo ID", "Alternate number 1", "Alternate number 2"}
_DATE_COLUMNS = {"ACCOUNT_ACTIVATION_DATE", "STATEMENT_DATE", "PREP_DATE", "PAYMENT_DUE_DATE", "LAST_PAYMENT_DT_TO_INVOICE",
                 "LAST_ADJ_DT_AGAINST_INVOICE", "LAST_DISCONNECTION_TSP_DATE"}


def _data_sheet(ws, columns: list[str], roles: dict[str, str]) -> None:
    ws.append(columns)
    seen = set()
    for i, (cell, name) in enumerate(zip(ws[1], columns), start=1):
        role = roles.get(name) if name not in seen else None  # only the first ACCOUNT_STATUS is read
        seen.add(name)
        cell.font = Font(bold=True, color="FFFFFF" if role else "1F2937")
        cell.fill = PatternFill("solid", fgColor=_ROLE_FILL[role] if role else "E5E7EB")
        cell.alignment = Alignment(vertical="center")
        letter = cell.column_letter
        ws.column_dimensions[letter].width = max(12, min(len(name) + 4, 40))
        fmt = "@" if name in _TEXT_COLUMNS else "DD-MM-YYYY" if name in _DATE_COLUMNS else None
        if fmt:
            for (c,) in ws.iter_rows(min_row=2, max_row=MAX_ROWS + 1, min_col=i, max_col=i):
                c.number_format = fmt
    ws.freeze_panes = "A2"


def template_workbook() -> bytes:
    wb = Workbook()
    _data_sheet(wb.active, INVOICE_COLUMNS, {c: r[0] for c, r in INVOICE_ROLES.items()})
    wb.active.title = "Invoice wise"
    _data_sheet(wb.create_sheet("Account Wise"), ACCOUNT_COLUMNS, {c: r[0] for c, r in ACCOUNT_ROLES.items()})

    notes = wb.create_sheet("Instructions")
    notes.append(["How to fill this template"])
    notes.append(["Same layout as the Smartflo collections export: one row per invoice on Invoice wise, one row per account on Account Wise (optional). "
                  "Keep every header as it is; columns not listed below can be left blank — they are stored with the call for reference."])
    notes.append([f"Up to {MAX_ROWS} invoice rows per upload. Rows are sorted by amount and can be edited before calls start."])
    notes.append([])
    notes.append(["Sheet", "Column", "Used for", "Required", "Example", "Notes"])
    for cell in notes[notes.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="002F51")
    labels = {"AI": "Sent to the AI agent", "Dial": "Dialling", "Skip": "Skip check", "Filter": "Filter"}
    entries = [("Invoice wise", col, role, required, example, note) for col, (role, required, example, note) in INVOICE_ROLES.items()]
    entries += [("Account Wise", col, role, False, "", note) for col, (role, note) in ACCOUNT_ROLES.items()]
    for sheet, col, role, required, example, note in entries:
        notes.append([sheet, col, labels[role], "Yes" if required else "", example, note])
        notes.cell(notes.max_row, 2).font = Font(bold=True)
        notes.cell(notes.max_row, 3).font = Font(color=_ROLE_FILL[role])
    notes.append([])
    notes.append(["Header colours: dark blue = sent to the AI, blue = dialling, orange = skip check, green = filter, grey = reference only."])
    notes["A1"].font = Font(bold=True, size=13, color="002F51")
    for col, width in zip("ABCDEF", (14, 28, 22, 10, 38, 90)):
        notes.column_dimensions[col].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
